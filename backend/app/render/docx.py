"""CV -> DOCX, mirroring the HTML layout with native Word constructs."""

from __future__ import annotations

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from ..models import CV
from .common import DOCX_FONTS, PAGE_SIZES_MM, prepared, safe_url


def _rgb(hex_color: str) -> RGBColor:
    h = hex_color.lstrip("#")
    if len(h) == 3:
        h = "".join(c * 2 for c in h)
    return RGBColor.from_string(h.upper())


def _set_font(style, name: str, size_pt: float) -> None:
    style.font.name = name
    style.font.size = Pt(size_pt)
    rfonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)


def _bottom_border(paragraph, hex_color: str) -> None:
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), str(_rgb(hex_color)))
    borders.append(bottom)
    paragraph._p.get_or_add_pPr().append(borders)


def _hyperlink(paragraph, text: str, url: str) -> None:
    rel_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    node = OxmlElement("w:t")
    node.text = text
    node.set(qn("xml:space"), "preserve")
    run.append(node)
    link.append(run)
    paragraph._p.append(link)


def render_docx(cv: CV) -> bytes:
    cv = prepared(cv)
    s = cv.style
    size = s.font_size_pt
    accent = _rgb(s.accent_color)
    page_w, page_h = PAGE_SIZES_MM[s.page_size]

    doc = Document()
    page = doc.sections[0]
    page.page_width, page.page_height = Mm(page_w), Mm(page_h)
    page.left_margin = page.right_margin = page.top_margin = page.bottom_margin = Mm(s.margin_mm)
    content_width = Mm(page_w - 2 * s.margin_mm)

    normal = doc.styles["Normal"]
    _set_font(normal, DOCX_FONTS[s.font], size)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = s.line_height
    _set_font(doc.styles["List Bullet"], DOCX_FONTS[s.font], size)

    body_align = WD_ALIGN_PARAGRAPH.JUSTIFY if s.text_align == "justify" else WD_ALIGN_PARAGRAPH.LEFT

    def para(space_before: float = 0, keep_next: bool = False, style: str | None = None):
        p = doc.add_paragraph(style=style)
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.keep_with_next = keep_next
        return p

    def row(left: str, right: str, bold: bool, space_before: float = 0):
        p = para(space_before, keep_next=True)
        p.paragraph_format.tab_stops.add_tab_stop(content_width, WD_TAB_ALIGNMENT.RIGHT)
        p.add_run(left).bold = bold
        if right:
            p.add_run("\t" + right).bold = bold
        return p

    def bullets(items) -> None:
        for b in items:
            p = para(style="List Bullet")
            p.paragraph_format.left_indent = Pt(size * 2.8)
            p.paragraph_format.first_line_indent = Pt(-size * 1.4)
            p.add_run(b.text)
            p.alignment = body_align

    align = WD_ALIGN_PARAGRAPH.CENTER if s.name_align == "center" else WD_ALIGN_PARAGRAPH.LEFT
    if cv.header.name:
        p = para()
        p.alignment = align
        run = p.add_run(cv.header.name)
        run.bold = True
        run.font.size = Pt(size * 2)
        run.font.color.rgb = accent
        p.paragraph_format.space_after = Pt(size * 0.35)
    if cv.header.headline:
        p = para()
        p.alignment = align
        p.add_run(cv.header.headline).bold = True
    if cv.header.contacts:
        p = para()
        p.alignment = align
        for i, c in enumerate(cv.header.contacts):
            if i:
                p.add_run(" | ")
            url = safe_url(c.url)
            if url:
                _hyperlink(p, c.text, url)
            else:
                p.add_run(c.text)

    for sec in cv.sections:
        titled = sec.show_title and bool(sec.title)
        gap = size * (1.1 if titled else 0.9)
        if titled:
            p = para(gap, keep_next=True)
            run = p.add_run(sec.title)
            run.bold = True
            run.font.size = Pt(size * 1.25)
            run.font.color.rgb = accent
            url = safe_url(sec.title_url)
            if url:
                p.add_run(" ")
                _hyperlink(p, f"[ {url.split('://', 1)[-1].rstrip('/')} ]", url)
            p.paragraph_format.space_after = Pt(size * 0.6)
            if s.heading_style == "rule":
                _bottom_border(p, s.accent_color)
            gap = 0
        for line in sec.text.split("\n") if sec.text else []:
            p = para(gap)
            p.add_run(line.strip())
            p.alignment = body_align
            gap = 0
        bullets(sec.bullets)
        for index, item in enumerate(sec.items):
            before = gap if index == 0 and not sec.bullets else size * 0.6 if index else 0
            gap = 0
            if item.title or item.date:
                row(item.title, item.date, True, before)
                before = 0
            if item.subtitle or item.location:
                row(item.subtitle, item.location, False, before)
                before = 0
            for line in item.lines:
                para(before).add_run(line.text)
                before = 0
            bullets(item.bullets)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
